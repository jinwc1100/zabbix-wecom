# -*- coding: utf-8 -*-
"""导出 Grafana 父/子系统及按 application.yml 匹配的责任人。"""
from __future__ import print_function

import ast
import http.cookiejar
import json
import os
import re
import sys
import urllib.request

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(SCRIPT_DIR)
YML_PATH = os.path.join(ROOT, "src", "main", "resources", "application.yml")
LIST_PATH = os.path.join(SCRIPT_DIR, "list.txt")
OUT_PATH = os.path.join(SCRIPT_DIR, "SLA父子系统责任人.xlsx")

SOURCES = [
    ("子系统", "afj5c5o5v3rb4f", ["子系统名称", "SLA名称", "服务名称", "子系统", "名称"]),
    ("父系统", "aftk5ryifcfeof", ["父系统名称", "SLA名称", "服务名称", "名称"]),
]
PANEL_IDS = (4, 2)  # 老版、新版


def load_config(path):
    text = open(path, encoding="utf-8").read()
    grafana_block = text[text.index("\ngrafana:"):text.index("\nwecom:", text.index("\ngrafana:"))]
    base = re.search(r"base-url:\s*(\S+)", grafana_block).group(1).rstrip("/")
    user = re.search(r"username:\s*(\S+)", grafana_block).group(1)
    password = re.search(r"password:\s*'([^']+)'", grafana_block).group(1)

    owner_block = text[text.index("owners:"):text.index("\nsla:", text.index("owners:"))]
    default_owner = re.search(r"default-owner:\s*\"([^\"]+)\"", owner_block).group(1)
    rules = []
    current = None
    for raw in owner_block.splitlines():
        line = raw.strip()
        name_match = re.match(r"- name:\s*\"([^\"]+)\"", line)
        if name_match:
            if current:
                rules.append(current)
            current = {"name": name_match.group(1), "mobile": "", "priority": 50,
                       "regions": [], "types": []}
            continue
        if current is None:
            continue
        value_match = re.match(r"(mobile|priority|regions|types):\s*(.*)", line)
        if not value_match:
            continue
        key, value = value_match.groups()
        if key == "mobile":
            current[key] = value.strip().strip("\"'")
        elif key == "priority":
            current[key] = int(value)
        else:
            current[key] = ast.literal_eval(value)
    if current:
        rules.append(current)
    return base, user, password, default_owner, rules


def load_mobiles_from_list(path):
    """list.txt 行尾「责任人：姓名 手机号」→ 姓名到手机。"""
    mobiles = {}
    for raw in open(path, encoding="utf-8"):
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        match = re.search(r"责任人[：:]\s*(\S+)\s+(\d{11})\s*$", line)
        if match:
            mobiles[match.group(1)] = match.group(2)
    return mobiles


def login(base, user, password):
    jar = http.cookiejar.CookieJar()
    opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(jar))
    request = urllib.request.Request(
        base + "/login",
        data=json.dumps({"user": user, "password": password}).encode("utf-8"),
        headers={"Content-Type": "application/json"},
    )
    opener.open(request, timeout=60)
    return opener


def fetch_dashboard(opener, base, uid):
    url = base + "/api/dashboards/uid/" + uid
    return json.loads(opener.open(url, timeout=60).read().decode("utf-8"))


def find_panel(dashboard, panel_id):
    pending = list(dashboard.get("dashboard", {}).get("panels") or [])
    while pending:
        panel = pending.pop(0)
        if panel.get("id") == panel_id:
            return panel
        pending.extend(panel.get("panels") or [])
    raise KeyError("panel %s not found" % panel_id)


def query_panel(opener, base, panel):
    target = panel["targets"][0]
    datasource = target["datasource"]
    query = {
        "refId": target.get("refId", "A"),
        "format": "table",
        "rawQuery": True,
        "rawSql": target["rawSql"],
        "datasource": {"type": datasource.get("type"), "uid": datasource.get("uid")},
    }
    payload = {"queries": [query], "from": "now-1h", "to": "now"}
    request = urllib.request.Request(
        base + "/api/ds/query",
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
    )
    response = json.loads(opener.open(request, timeout=300).read().decode("utf-8"))
    result = (response.get("results") or {}).get(query["refId"]) or {}
    if result.get("error"):
        raise RuntimeError(result["error"])
    rows = []
    for frame in result.get("frames") or []:
        fields = [field.get("name") for field in frame["schema"]["fields"]]
        values = frame.get("data", {}).get("values") or []
        count = len(values[0]) if values else 0
        for index in range(count):
            rows.append(dict((name, values[column][index])
                             for column, name in enumerate(fields)))
    return rows


def detect_type(system_name):
    if any(token in system_name for token in ("开评标", "开标", "评标", "编标")):
        return "BID"
    if any(token in system_name for token in ("交易系统", "交易平台", "交易")):
        return "TRADE"
    return "NONE"


def region_hit(system_name, regions):
    return any(region and region.strip() in system_name for region in regions)


def pick_best(rules):
    return sorted(rules, key=lambda rule: rule["priority"], reverse=True)[0]


def resolve_owner(system_name, rules, default_owner):
    system_type = detect_type(system_name)
    typed = []
    region_only = []
    for rule in rules:
        if not region_hit(system_name, rule["regions"]):
            continue
        if rule["types"]:
            if system_type != "NONE" and system_type in rule["types"]:
                typed.append(rule)
        else:
            region_only.append(rule)
    if typed:
        owner = pick_best(typed)
        return owner["name"]
    if region_only:
        owner = pick_best(region_only)
        return owner["name"]

    fallback = []
    for rule in rules:
        if not region_hit(system_name, rule["regions"]):
            continue
        if not rule["types"] or "TRADE" in rule["types"]:
            fallback.append(rule)
    if fallback:
        return pick_best(fallback)["name"]
    return default_owner


def first_name(row, candidates):
    for column in candidates:
        value = row.get(column)
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def collect_systems(opener, base):
    result = []
    seen = set()
    for level, uid, candidates in SOURCES:
        dashboard = fetch_dashboard(opener, base, uid)
        for panel_id in PANEL_IDS:
            rows = query_panel(opener, base, find_panel(dashboard, panel_id))
            print("%s panel=%s rows=%d" % (level, panel_id, len(rows)))
            for row in rows:
                name = first_name(row, candidates)
                key = (level, name)
                if name and key not in seen:
                    seen.add(key)
                    result.append(key)
    return result


def write_excel(systems, rules, default_owner, mobiles):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "父子系统责任人"
    headers = ["子系统/父系统", "系统名称", "责任人", "手机号码"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    for level, name in sorted(systems, key=lambda item: (item[0], item[1])):
        owner = resolve_owner(name, rules, default_owner)
        sheet.append([level, name, owner, mobiles.get(owner, "")])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in zip(("A", "B", "C", "D"), (16, 52, 16, 20)):
        sheet.column_dimensions[column].width = width
    workbook.save(OUT_PATH)


def main():
    base, user, password, default_owner, rules = load_config(YML_PATH)
    mobiles = load_mobiles_from_list(LIST_PATH)
    print("list.txt 责任人手机 %d 人" % len(mobiles))
    print("登录 Grafana: %s" % base)
    opener = login(base, user, password)
    systems = collect_systems(opener, base)
    write_excel(systems, rules, default_owner, mobiles)
    print("系统数=%d，已写入: %s" % (len(systems), OUT_PATH))


if __name__ == "__main__":
    main()
